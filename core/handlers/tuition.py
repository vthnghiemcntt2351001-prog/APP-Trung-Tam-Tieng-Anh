from decimal import Decimal, InvalidOperation
from datetime import date

from django.utils.dateparse import parse_date
from django.db.models import Q, Sum
from django.contrib import messages
from django.shortcuts import render, redirect, get_object_or_404

from ..models import Student, EnglishClass, Tuition, PaymentTransaction, Enrollment
from ..permissions import admin_required, finance_required, admin_context


@finance_required
def tuition_list(request):
    tuitions = Tuition.objects.select_related(
        "student",
        "english_class",
        "english_class__course",
    ).prefetch_related("transactions").all()

    keyword = request.GET.get("q", "").strip()

    if keyword:
        tuitions = tuitions.filter(
            Q(student__full_name__icontains=keyword)
            | Q(english_class__name__icontains=keyword)
        )

    total_fee = (
        tuitions.aggregate(total=Sum("total_fee"))["total"]
        or Decimal("0")
    )

    total_paid = (
        tuitions.aggregate(total=Sum("paid_amount"))["total"]
        or Decimal("0")
    )

    total_debt = sum(
        (item.remaining_amount for item in tuitions),
        Decimal("0"),
    )

    from django.core.paginator import Paginator
    paginator = Paginator(tuitions, 15)
    page_number = request.GET.get("page")
    page_obj = paginator.get_page(page_number)

    context = {
        "page_obj": page_obj,
        "tuitions": page_obj,
        "keyword": keyword,
        "total_fee": total_fee,
        "total_paid": total_paid,
        "total_debt": total_debt,
        **admin_context(request),
    }

    return render(request, "core/finance/tuition/list.html", context)


@finance_required
def tuition_add(request):
    import json
    from collections import defaultdict

    students = Student.objects.filter(is_active=True).order_by("full_name")
    classes = EnglishClass.objects.select_related("course").filter(is_active=True).order_by("name")

    student_classes_map = defaultdict(list)
    enrollments = (
        Enrollment.objects
        .filter(status=Enrollment.Status.STUDYING)
        .select_related("english_class", "english_class__course")
    )
    for en in enrollments:
        if en.english_class and en.english_class.course:
            student_classes_map[str(en.student_id)].append({
                "id": en.english_class_id,
                "name": en.english_class.name,
                "course_name": en.english_class.course.name,
                "tuition": float(en.english_class.course.tuition),
            })

    if request.method == "POST":
        student_id = request.POST.get("student")
        class_id = request.POST.get("english_class")
        paid_amount_raw = request.POST.get(
            "paid_amount", "0"
        ).strip()
        payment_method = request.POST.get(
            "payment_method", PaymentTransaction.PaymentMethod.CASH
        )
        payment_date_raw = request.POST.get("payment_date") or None
        note = request.POST.get("note", "").strip()

        try:
            student = Student.objects.get(pk=student_id)

            english_class = EnglishClass.objects.select_related(
                "course"
            ).get(pk=class_id)

            if not english_class.students.filter(pk=student.pk).exists():
                messages.error(
                    request,
                    "Học viên chưa đăng ký vào lớp này."
                )
                return redirect("tuition_add")

            paid_amount = Decimal(paid_amount_raw)

            if not paid_amount.is_finite() or paid_amount < 0:
                raise ValueError

            if paid_amount > english_class.course.tuition:
                messages.error(
                    request,
                    "Số tiền đã đóng không được lớn hơn học phí khóa học."
                )
                return redirect("tuition_add")

            payment_date = (
                parse_date(payment_date_raw)
                if payment_date_raw else None
            )

            if payment_date_raw and not payment_date:
                messages.error(
                    request,
                    "Ngày thanh toán không hợp lệ."
                )
                return redirect("tuition_add")

            # Kiểm tra nếu học viên đã có bản ghi Tuition cho lớp này
            existing_tuition = Tuition.objects.filter(
                student=student,
                english_class=english_class,
            ).first()

            if existing_tuition:
                tuition = existing_tuition
                if tuition.remaining_amount <= 0:
                    messages.info(
                        request,
                        f"Học viên '{student.full_name}' đã hoàn tất 100% học phí cho lớp '{english_class.name}'."
                    )
                    return redirect("tuition_list")
            else:
                tuition = Tuition.objects.create(
                    student=student,
                    english_class=english_class,
                    total_fee=english_class.course.tuition,
                    paid_amount=0,
                    note=note,
                )

            if paid_amount > 0:
                PaymentTransaction.objects.create(
                    tuition=tuition,
                    amount=paid_amount,
                    payment_method=payment_method,
                    payment_date=payment_date or date.today(),
                    collected_by=request.user,
                    note=note or "Thanh toán học phí",
                )

            messages.success(
                request,
                "Tạo phiếu học phí thành công!"
            )
            return redirect("tuition_list")

        except (Student.DoesNotExist, EnglishClass.DoesNotExist):
            messages.error(
                request,
                "Học viên hoặc lớp học không tồn tại."
            )

        except (InvalidOperation, ValueError, TypeError):
            messages.error(
                request,
                "Số tiền đã đóng không hợp lệ."
            )

    return render(request, "core/finance/tuition/add.html", {
        "students": students,
        "classes": classes,
        "student_classes_json": json.dumps(student_classes_map),
        "payment_methods": PaymentTransaction.PaymentMethod.choices,
        "today": date.today().isoformat(),
        **admin_context(request),
    })


@finance_required
def tuition_payment_add(request, tuition_id):
    """Ghi nhận đợt thanh toán học phí mới (trả góp / đóng đợt tiếp theo)."""
    tuition = get_object_or_404(
        Tuition.objects.select_related(
            "student",
            "english_class",
            "english_class__course",
        ),
        pk=tuition_id,
    )

    if request.method == "POST":
        amount_raw = request.POST.get("amount", "0").strip()
        payment_method = request.POST.get(
            "payment_method", PaymentTransaction.PaymentMethod.CASH
        )
        payment_date_raw = request.POST.get("payment_date") or None
        note = request.POST.get("note", "").strip()

        try:
            amount = Decimal(amount_raw)
            if not amount.is_finite() or amount <= 0:
                messages.error(request, "Số tiền nộp phải lớn hơn 0 VNĐ.")
                return redirect("tuition_payment_add", tuition_id=tuition_id)

            if amount > tuition.remaining_amount:
                messages.error(
                    request,
                    f"Số tiền nộp ({amount:,.0f} đ) không được lớn hơn số tiền còn nợ ({tuition.remaining_amount:,.0f} đ)."
                )
                return redirect("tuition_payment_add", tuition_id=tuition_id)

            payment_date = parse_date(payment_date_raw) if payment_date_raw else date.today()
            if not payment_date:
                messages.error(request, "Ngày thanh toán không hợp lệ.")
                return redirect("tuition_payment_add", tuition_id=tuition_id)

            PaymentTransaction.objects.create(
                tuition=tuition,
                amount=amount,
                payment_method=payment_method,
                payment_date=payment_date,
                collected_by=request.user,
                note=note or "Nộp học phí đợt tiếp theo",
            )

            messages.success(
                request,
                f"Thu {amount:,.0f} VNĐ cho học viên '{tuition.student.full_name}' thành công!"
            )
            return redirect("tuition_list")

        except (InvalidOperation, ValueError, TypeError) as exc:
            messages.error(request, f"Dữ liệu thanh toán không hợp lệ: {exc}")

    return render(request, "core/finance/tuition/payment_add.html", {
        "tuition": tuition,
        "payment_methods": PaymentTransaction.PaymentMethod.choices,
        "today": date.today().isoformat(),
        **admin_context(request),
    })



@finance_required
def tuition_edit(request, tuition_id):
    tuition = get_object_or_404(
        Tuition.objects.select_related(
            "student",
            "english_class__course",
        ),
        pk=tuition_id,
    )

    if request.method == "POST":
        paid_amount_raw = request.POST.get(
            "paid_amount", ""
        ).strip()
        payment_date_raw = request.POST.get("payment_date") or None
        note = request.POST.get("note", "").strip()

        try:
            paid_amount = Decimal(paid_amount_raw)

            if not paid_amount.is_finite() or paid_amount < 0:
                raise ValueError

            if paid_amount > tuition.total_fee:
                messages.error(
                    request,
                    "Số tiền đã đóng không được lớn hơn tổng học phí."
                )
                return redirect(
                    "tuition_edit",
                    tuition_id=tuition.id,
                )

            payment_date = (
                parse_date(payment_date_raw)
                if payment_date_raw else None
            )

            if payment_date_raw and not payment_date:
                messages.error(
                    request,
                    "Ngày thanh toán không hợp lệ."
                )
                return redirect(
                    "tuition_edit",
                    tuition_id=tuition.id,
                )

            tuition.paid_amount = paid_amount
            tuition.payment_date = payment_date
            tuition.note = note
            tuition.save()

            messages.success(
                request,
                "Cập nhật phiếu học phí thành công!"
            )
            return redirect("tuition_list")

        except (InvalidOperation, ValueError, TypeError):
            messages.error(
                request,
                "Số tiền đã đóng không hợp lệ."
            )

    return render(request, "core/finance/tuition/edit.html", {
        "tuition": tuition,
        "today": date.today().isoformat(),
        **admin_context(request),
    })


@admin_required
def tuition_delete(request, tuition_id):
    if request.method != "POST":
        return redirect("tuition_list")

    tuition = get_object_or_404(Tuition, pk=tuition_id)
    tuition.delete()

    messages.success(request, "Xóa phiếu học phí thành công!")
    return redirect("tuition_list")


@finance_required
def tuition_receipt(request, tuition_id):
    from ..utils import num_to_vietnamese_words

    tuition = get_object_or_404(
        Tuition.objects.select_related(
            "student",
            "english_class",
            "english_class__course",
            "english_class__teacher",
        ).prefetch_related("transactions"),
        pk=tuition_id,
    )

    transaction_id = request.GET.get("transaction_id")
    selected_transaction = None
    if transaction_id:
        selected_transaction = tuition.transactions.filter(pk=transaction_id).first()
    if not selected_transaction:
        selected_transaction = tuition.transactions.order_by("-payment_date", "-id").first()

    receipt_amount = selected_transaction.amount if selected_transaction else tuition.paid_amount
    amount_in_words = num_to_vietnamese_words(receipt_amount)
    receipt_code = selected_transaction.receipt_code if selected_transaction else f"PT-{tuition.id:05d}"
    receipt_date = selected_transaction.payment_date if selected_transaction else (tuition.payment_date or date.today())
    payment_method_label = (
        selected_transaction.get_payment_method_display()
        if selected_transaction
        else "Tiền mặt / Chuyển khoản"
    )

    context = {
        "tuition": tuition,
        "selected_transaction": selected_transaction,
        "transactions": tuition.transactions.all(),
        "amount_in_words": amount_in_words,
        "receipt_amount": receipt_amount,
        "receipt_code": receipt_code,
        "receipt_date": receipt_date,
        "payment_method_label": payment_method_label,
        "today": date.today(),
        **admin_context(request),
    }
    return render(request, "core/finance/tuition/receipt.html", context)



@finance_required
def tuition_export_excel(request):
    import io
    import datetime
    from django.http import HttpResponse
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    keyword = request.GET.get("q", "").strip()
    tuitions = Tuition.objects.select_related(
        "student",
        "english_class",
        "english_class__course",
    ).all()

    if keyword:
        tuitions = tuitions.filter(
            Q(student__full_name__icontains=keyword)
            | Q(english_class__name__icontains=keyword)
        )

    tuitions = tuitions.order_by("id")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Báo cáo học phí"
    ws.views.sheetView[0].showGridLines = True

    title_font = Font(name="Arial", size=15, bold=True, color="1E3A8A")
    meta_font = Font(name="Arial", size=10, italic=True, color="475569")
    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="203864", end_color="203864", fill_type="solid")
    summary_font = Font(name="Arial", size=10, bold=True)
    summary_fill = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")
    data_font = Font(name="Arial", size=10)
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")
    right_align = Alignment(horizontal="right", vertical="center")

    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    # Tiêu đề
    ws.merge_cells("A1:J1")
    ws["A1"] = "BÁO CÁO THEO DÕI HỌC PHÍ & CÔNG NỢ - ENGLISH CENTER"
    ws["A1"].font = title_font
    ws["A1"].alignment = center_align
    ws.row_dimensions[1].height = 32

    ws.merge_cells("A2:J2")
    ws["A2"] = f"Ngày xuất báo cáo: {datetime.date.today().strftime('%d/%m/%Y')} | Tổng số phiếu: {tuitions.count()}"
    ws["A2"].font = meta_font
    ws["A2"].alignment = center_align
    ws.row_dimensions[2].height = 18

    # Hàng tiêu đề bảng
    headers = [
        "STT", "Mã phiếu", "Học viên", "Lớp học", "Khóa học",
        "Tổng học phí (VNĐ)", "Đã đóng (VNĐ)", "Còn nợ (VNĐ)", "Trạng thái", "Ngày thanh toán"
    ]
    ws.row_dimensions[4].height = 26

    for col_num, header_title in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col_num, value=header_title)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = thin_border

    # Hàng dữ liệu
    row_num = 5
    total_fee_sum = 0
    total_paid_sum = 0
    total_debt_sum = 0

    for idx, item in enumerate(tuitions, 1):
        ws.row_dimensions[row_num].height = 22
        pay_date_str = item.payment_date.strftime("%d/%m/%Y") if item.payment_date else "Chưa thanh toán"
        fee = float(item.total_fee)
        paid = float(item.paid_amount)
        debt = float(item.remaining_amount)

        total_fee_sum += fee
        total_paid_sum += paid
        total_debt_sum += debt

        row_data = [
            (idx, center_align, None),
            (f"PT-{item.id:05d}", center_align, None),
            (item.student.full_name, left_align, None),
            (item.english_class.name, left_align, None),
            (item.english_class.course.name if item.english_class.course else "-", left_align, None),
            (fee, right_align, '#,##0'),
            (paid, right_align, '#,##0'),
            (debt, right_align, '#,##0'),
            (item.payment_status, center_align, None),
            (pay_date_str, center_align, None),
        ]

        for col_num, (val, align, num_fmt) in enumerate(row_data, 1):
            cell = ws.cell(row=row_num, column=col_num, value=val)
            cell.font = data_font
            cell.alignment = align
            cell.border = thin_border
            if num_fmt:
                cell.number_format = num_fmt
            if idx % 2 == 0:
                cell.fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

        row_num += 1

    # Dòng tổng cộng
    ws.row_dimensions[row_num].height = 24
    ws.merge_cells(f"A{row_num}:E{row_num}")
    sum_label_cell = ws.cell(row=row_num, column=1, value="TỔNG CỘNG TOÀN BỘ")
    sum_label_cell.font = summary_font
    sum_label_cell.alignment = Alignment(horizontal="right", vertical="center")
    sum_label_cell.fill = summary_fill

    for c in range(1, 6):
        ws.cell(row=row_num, column=c).fill = summary_fill
        ws.cell(row=row_num, column=c).border = thin_border

    fee_sum_cell = ws.cell(row=row_num, column=6, value=total_fee_sum)
    fee_sum_cell.font = summary_font
    fee_sum_cell.fill = summary_fill
    fee_sum_cell.alignment = right_align
    fee_sum_cell.number_format = '#,##0'
    fee_sum_cell.border = thin_border

    paid_sum_cell = ws.cell(row=row_num, column=7, value=total_paid_sum)
    paid_sum_cell.font = summary_font
    paid_sum_cell.fill = summary_fill
    paid_sum_cell.alignment = right_align
    paid_sum_cell.number_format = '#,##0'
    paid_sum_cell.border = thin_border

    debt_sum_cell = ws.cell(row=row_num, column=8, value=total_debt_sum)
    debt_sum_cell.font = summary_font
    debt_sum_cell.fill = summary_fill
    debt_sum_cell.alignment = right_align
    debt_sum_cell.number_format = '#,##0'
    debt_sum_cell.border = thin_border

    for c in range(9, 11):
        empty_sum = ws.cell(row=row_num, column=c, value="")
        empty_sum.fill = summary_fill
        empty_sum.border = thin_border

    # Căn chỉnh độ rộng cột
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = max(len(str(cell.value or '')) for cell in col)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"Bao_Cao_Hoc_Phi_{datetime.date.today().strftime('%Y%m%d')}.xlsx"
    response = HttpResponse(
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response

