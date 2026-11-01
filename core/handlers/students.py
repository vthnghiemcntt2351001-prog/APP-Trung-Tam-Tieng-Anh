from django.db.models import Q
from django.contrib import messages
from django.shortcuts import render, redirect, get_object_or_404

from ..models import Student
from ..permissions import admin_required, academic_required, admin_context


@academic_required
def student_list(request):
    keyword = request.GET.get("q", "").strip()
    status_filter = request.GET.get("status", "active")

    students = Student.objects.all()

    if status_filter == "archived":
        students = students.filter(is_active=False)
    elif status_filter == "all":
        pass
    else:
        status_filter = "active"
        students = students.filter(is_active=True)

    if keyword:
        students = students.filter(
            Q(full_name__icontains=keyword)
            | Q(email__icontains=keyword)
            | Q(phone__icontains=keyword)
        )

    from django.core.paginator import Paginator

    students = students.order_by("-id")

    paginator = Paginator(students, 15)
    page_number = request.GET.get("page")
    page_obj = paginator.get_page(page_number)

    total_active = Student.objects.filter(is_active=True).count()
    total_archived = Student.objects.filter(is_active=False).count()

    context = {
        "page_obj": page_obj,
        "students": page_obj,
        "keyword": keyword,
        "status_filter": status_filter,
        "total_active": total_active,
        "total_archived": total_archived,
        **admin_context(request),
    }

    return render(request, "core/students/list.html", context)


@academic_required
def student_add(request):
    from ..forms import StudentForm
    form = StudentForm(request.POST or None)

    if request.method == "POST":
        if form.is_valid():
            student = form.save()
            messages.success(request, f"Thêm học viên '{student.full_name}' thành công!")
            return redirect("student_list")
        else:
            for field, errors in form.errors.items():
                for err in errors:
                    messages.error(request, err)

    return render(
        request,
        "core/students/add.html",
        {
            "form": form,
            **admin_context(request),
        },
    )


@academic_required
def student_edit(request, student_id):
    student = get_object_or_404(Student, id=student_id)
    from ..forms import StudentForm
    form = StudentForm(request.POST or None, instance=student)

    if request.method == "POST":
        if form.is_valid():
            form.save()
            messages.success(
                request,
                f"Cập nhật học viên '{student.full_name}' thành công!"
            )
            return redirect("student_list")
        else:
            for field, errors in form.errors.items():
                for err in errors:
                    messages.error(request, err)

    return render(
        request,
        "core/students/edit.html",
        {
            "student": student,
            "form": form,
            **admin_context(request),
        },
    )

    context = {
        "student": student,
        **admin_context(request),
    }

    return render(request, "core/students/edit.html", context)


@admin_required
def student_delete(request, student_id):
    student = get_object_or_404(Student, id=student_id)

    if request.method == "POST":
        action = request.POST.get("action", "archive")
        if action == "restore":
            student.is_active = True
            student.save(update_fields=["is_active"])
            messages.success(request, f"Đã khôi phục học viên '{student.full_name}' hoạt động trở lại.")
        else:
            student.is_active = False
            student.save(update_fields=["is_active"])
            messages.success(request, f"Đã lưu trữ học viên '{student.full_name}'. Toàn bộ lịch sử học phí và điểm số vẫn được bảo lưu an toàn.")
        return redirect("student_list")

    return render(request, "core/students/delete.html", {
        "student": student,
        **admin_context(request),
    })


@academic_required
def student_export_excel(request):
    import io
    import datetime
    from django.http import HttpResponse
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    keyword = request.GET.get("q", "").strip()
    students = Student.objects.all()

    if keyword:
        students = students.filter(
            Q(full_name__icontains=keyword)
            | Q(email__icontains=keyword)
            | Q(phone__icontains=keyword)
        )

    students = students.order_by("id")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Danh sách học viên"
    ws.views.sheetView[0].showGridLines = True

    title_font = Font(name="Arial", size=15, bold=True, color="1E3A8A")
    meta_font = Font(name="Arial", size=10, italic=True, color="475569")
    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="203864", end_color="203864", fill_type="solid")
    data_font = Font(name="Arial", size=10)
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")

    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    # Tiêu đề
    ws.merge_cells("A1:G1")
    ws["A1"] = "DANH SÁCH HỌC VIÊN - ENGLISH CENTER"
    ws["A1"].font = title_font
    ws["A1"].alignment = center_align
    ws.row_dimensions[1].height = 32

    ws.merge_cells("A2:G2")
    ws["A2"] = f"Ngày xuất báo cáo: {datetime.date.today().strftime('%d/%m/%Y')} | Tổng số hồ sơ: {students.count()}"
    ws["A2"].font = meta_font
    ws["A2"].alignment = center_align
    ws.row_dimensions[2].height = 18

    # Hàng tiêu đề bảng
    headers = ["STT", "Mã HV", "Họ và tên", "Email", "Số điện thoại", "Ngày sinh", "Địa chỉ"]
    ws.row_dimensions[4].height = 26

    for col_num, header_title in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col_num, value=header_title)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = thin_border

    # Hàng dữ liệu
    row_num = 5
    for idx, st in enumerate(students, 1):
        ws.row_dimensions[row_num].height = 22
        dob_str = st.date_of_birth.strftime("%d/%m/%Y") if st.date_of_birth else "-"

        row_data = [
            (idx, center_align),
            (f"HV-{st.id:04d}", center_align),
            (st.full_name, left_align),
            (st.email, left_align),
            (st.phone, center_align),
            (dob_str, center_align),
            (st.address or "-", left_align),
        ]

        for col_num, (val, align) in enumerate(row_data, 1):
            cell = ws.cell(row=row_num, column=col_num, value=val)
            cell.font = data_font
            cell.alignment = align
            cell.border = thin_border
            if idx % 2 == 0:
                cell.fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

        row_num += 1

    # Căn chỉnh độ rộng cột
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = max(len(str(cell.value or '')) for cell in col)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"Danh_Sach_Hoc_Vien_{datetime.date.today().strftime('%Y%m%d')}.xlsx"
    response = HttpResponse(
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response

